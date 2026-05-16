from __future__ import annotations

import io
import re

from app.services.review_templates import ReviewTemplate


def review_to_docx_bytes(
    review: str,
    topic: str,
    citations: list[dict],
    *,
    template: ReviewTemplate = "literature_review",
    lang: str = "zh",
) -> bytes:
    try:
        from docx import Document
    except ImportError as e:
        raise RuntimeError(
            "DOCX 导出需要 python-docx：cd apps/api && .venv/bin/pip install python-docx"
        ) from e

    doc = Document()
    title = topic.strip() or ("Literature Review" if lang == "en" else "文献综述")
    doc.add_heading(title, level=0)
    template_labels = {
        "grant_proposal": ("模板：项目申请书（研究现状）", "Template: Grant proposal"),
        "quick_review": ("模板：文献速览", "Template: Quick review"),
        "comparative_review": ("模板：对比综述", "Template: Comparative review"),
    }
    if template in template_labels:
        zh_l, en_l = template_labels[template]
        doc.add_paragraph(zh_l if lang == "zh" else en_l)

    for para in re.split(r"\n{2,}", review.strip()):
        p = para.strip()
        if not p:
            continue
        if p.startswith("#"):
            level = min(3, len(p) - len(p.lstrip("#")))
            doc.add_heading(p.lstrip("#").strip(), level=level)
        else:
            doc.add_paragraph(p)

    if citations:
        doc.add_heading("References" if lang == "en" else "引用来源", level=1)
        for c in citations:
            ref = c.get("ref")
            title_s = c.get("title") or c.get("paper_id")
            section = c.get("section_path") or ""
            line = f"[{ref}] {title_s}"
            if section:
                line += f" — {section}"
            doc.add_paragraph(line, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
